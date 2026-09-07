from __future__ import annotations

import copy
import json
import re
import shutil
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.shared import Cm, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
DOC_DIR = ROOT / "doc"
PRD = DOC_DIR / "DeepSeek大模型金融投研系统-PRD.md"
WORK = DOC_DIR / ".srs-work"
UML_DIR = DOC_DIR / "uml"
OUT = DOC_DIR / "01_JARVIS金融投研平台_Software Requirement Specification_V1.0.docx"
SOURCE_MD = WORK / "JARVIS金融投研平台-SRS源.md"
FUNCTION_DATA = WORK / "function_data.json"
PLANTUML = DOC_DIR / "Skills" / "uml-generator" / "plantuml.jar"
GENERATOR = DOC_DIR / "Skills" / "md-to-srs-docx" / "scripts" / "generate_srs.py"
STYLES = DOC_DIR / "Skills" / "md-to-srs-docx" / "assets" / "template_styles.json"


DIAGRAMS = {
    "用例图_总览.puml": r'''@startuml
left to right direction
skinparam packageStyle rectangle
skinparam actorStyle awesome
skinparam shadowing false
skinparam dpi 140
skinparam usecaseBackgroundColor #F1F8E9
skinparam usecaseBorderColor #689F38
skinparam rectangleBackgroundColor #E3F2FD
skinparam rectangleBorderColor #1565C0
actor "普通用户" as User
actor "管理员" as Admin
actor "GitHub / Resend" as AuthProvider
actor "行情数据源" as MarketProvider
actor "AI 模型服务" as AiProvider
rectangle "JARVIS 金融投研平台" {
  usecase "身份登录与注册" as UC1
  usecase "查看多市场行情" as UC2
  usecase "AI 投研对话" as UC3
  usecase "模拟交易" as UC4
  usecase "策略回测" as UC5
  usecase "财报解析" as UC6
  usecase "询报价与趋势" as UC7
  usecase "研报情感分析" as UC8
  usecase "产业链挖掘" as UC9
  usecase "风险预警" as UC10
  usecase "账户管理" as UC11
  usecase "AI 配额管理" as UC12
  usecase "功能权限与审计" as UC13
  usecase "健康检查与运维" as UC14
  UC1 -[hidden]right- UC2
  UC2 -[hidden]right- UC3
  UC3 -[hidden]right- UC4
  UC4 -[hidden]right- UC5
  UC6 -[hidden]right- UC7
  UC7 -[hidden]right- UC8
  UC8 -[hidden]right- UC9
  UC9 -[hidden]right- UC10
  UC11 -[hidden]right- UC12
  UC12 -[hidden]right- UC13
  UC13 -[hidden]right- UC14
  UC1 -[hidden]down- UC6
  UC6 -[hidden]down- UC11
}
User --> UC1
User --> UC2
User --> UC3
User --> UC4
User --> UC5
User --> UC6
User --> UC7
User --> UC8
User --> UC9
User --> UC10
Admin --> UC11
Admin --> UC12
Admin --> UC13
Admin --> UC14
AuthProvider --> UC1
MarketProvider --> UC2
AiProvider --> UC3
AiProvider --> UC6
AiProvider --> UC7
AiProvider --> UC8
AiProvider --> UC9
@enduml
''',
    "流程图_主业务流程.puml": r'''@startuml
skinparam shadowing false
skinparam dpi 140
|用户|
start
:打开 JARVIS 工作台;
if (是否已登录?) then (否)
  :选择邮箱或 GitHub 登录;
  |Java 业务后端|
  :校验 CSRF、验证码或 OAuth state;
  if (认证通过?) then (否)
    :返回安全错误提示并限流;
    stop
  else (是)
    :签发 HttpOnly JWT 会话;
  endif
else (是)
endif
|用户|
:选择市场与标的;
|Java 业务后端|
:读取行情快照与 K 线;
if (数据源可用?) then (否)
  :使用缓存并标记数据时效;
else (是)
  :计算指标并返回行情;
endif
|用户|
:发起 AI 投研请求;
|Java 业务后端|
:校验功能权限与 AI 配额;
if (允许调用?) then (否)
  :返回额度不足或无权限;
  stop
else (是)
  :携带内部令牌调用 Python AI 服务;
endif
|Python AI 服务|
:完成数值处理与模型推理;
:返回结构化结果或流式文本;
|Java 业务后端|
:记录用量、审计并返回结果;
|用户|
:查看分析结果与风险提示;
stop
@enduml
''',
    "用例图_用户身份与登录管理（GitHub与邮箱）.puml": r'''@startuml
left to right direction
skinparam shadowing false
skinparam dpi 140
actor "用户" as U
actor "GitHub" as G
actor "Resend" as R
rectangle "身份与登录管理" {
  usecase "邮箱注册" as A
  usecase "邮箱验证码校验" as B
  usecase "邮箱登录" as C
  usecase "GitHub OAuth 登录" as D
  usecase "会话保持与登出" as E
  usecase "找回密码" as F
}
U --> A
U --> C
U --> D
U --> E
U --> F
A ..> B : <<include>>
F ..> B : <<include>>
G --> D
R --> B
@enduml
''',
    "用例图_行情中心（实时行情与 K 线）.puml": r'''@startuml
left to right direction
skinparam shadowing false
skinparam dpi 140
actor "普通用户" as U
actor "行情数据源" as M
rectangle "行情中心" {
  usecase "选择市场与标的" as A
  usecase "查看实时报价" as B
  usecase "查看 K 线" as C
  usecase "查看技术指标" as D
  usecase "获取缓存兜底" as E
}
U --> A
U --> B
U --> C
U --> D
B ..> E : <<extend>>
C ..> E : <<extend>>
M --> B
M --> C
@enduml
''',
    "用例图_AI 智能对话（投研助手）.puml": r'''@startuml
left to right direction
skinparam shadowing false
skinparam dpi 140
actor "普通用户" as U
actor "Java 业务后端" as J
actor "Python AI 服务" as P
actor "AI 模型服务" as A
rectangle "AI 投研助手" {
  usecase "发送问题" as C
  usecase "流式返回" as S
  usecase "校验权限与配额" as Q
  usecase "生成风险提示" as R
}
U --> C
C ..> Q : <<include>>
C ..> S : <<include>>
C ..> R : <<include>>
J --> Q
J --> P
P --> A
@enduml
''',
    "用例图_模拟盘交易.puml": r'''@startuml
left to right direction
skinparam shadowing false
skinparam dpi 140
actor "普通用户" as U
actor "行情服务" as M
rectangle "模拟盘交易" {
  usecase "提交模拟订单" as A
  usecase "校验余额与保证金" as B
  usecase "管理持仓" as C
  usecase "触发强平" as D
  usecase "查看成交记录" as E
}
U --> A
U --> C
U --> E
A ..> B : <<include>>
B ..> M : <<include>>
D ..> M : <<include>>
@enduml
''',
    "架构图_系统总体架构.puml": r'''@startuml
title JARVIS 金融投研平台系统总体架构
left to right direction
skinparam componentStyle rectangle
skinparam shadowing false
skinparam dpi 150
skinparam defaultFontName Microsoft YaHei
skinparam packageStyle rectangle
skinparam packageBackgroundColor #F7FAFF
skinparam packageBorderColor #5B7DB1
skinparam componentBackgroundColor #EAF2FF
skinparam componentBorderColor #3568A8
skinparam databaseBackgroundColor #FFF6D8
skinparam databaseBorderColor #C78A00
skinparam cloudBackgroundColor #EEF7EE
skinparam cloudBorderColor #4F8A4F

package "访问层" {
  component "Vue 3 + Vite 前端" as FE
  component "ECharts 行情与指标可视化" as Charts
  component "React Native 移动端（预留）" as Mobile
}

package "Java 业务与 CRUD 层  :8200" {
  component "Spring Boot 3.3 API" as Java
  component "Spring Security + JWT + CSRF" as Security
  component "账户 / RBAC / 配额 / 审计" as Account
  component "行情 / K 线 / 指标" as Market
  component "回测 / 模拟盘 / 风控" as Trading
  component "AI 网关与 SSE" as Gateway
  component "WebClient 定时采集" as Collector
}

package "AI 内部服务层  :8100" {
  component "Python FastAPI" as Py
  component "数值计算与数据预处理" as Calc
  component "DeepSeek 兼容调用适配器" as Adapter
}

database "PostgreSQL 生产\nH2 本地开发" as DB
component "Flyway 数据库迁移" as Flyway
cloud "DeepSeek / AI Provider" as LLM
cloud "GitHub OAuth" as Github
cloud "Resend 邮件验证码" as Resend
cloud "行情数据源" as Data
component "Nginx / Cloudflare HTTPS 入口" as Edge

FE --> Edge : HTTPS /api
Mobile --> Edge : HTTPS /api
Edge --> Java : 反向代理
FE --> Charts
Java --> Security
Java --> Account
Java --> Market
Java --> Trading
Java --> Gateway
Java --> Collector
Java --> DB : JPA / 事务
Flyway --> DB
Gateway --> Py : 内部令牌
Py --> Calc
Py --> Adapter
Adapter --> LLM : 服务端密钥
Security --> Github : OAuth state
Account --> Resend : 服务端 API Key
Collector --> Data : HTTPS
Market --> Data : 备用源 / 缓存
note right of Py
浏览器不直接访问 Python
AI Key 仅存在服务端
数值计算由 Python 完成
end note
@enduml
''',
}


FUNCTIONS = [
    ("F001", "用户身份与登录管理（GitHub与邮箱）", [
        ("邮箱注册与验证码", "用户提交邮箱、密码和一次性验证码；系统校验有效期、单次使用和错误次数后创建账户。", "验证输入和发送频率，调用 Resend 发送验证码，服务端哈希存储验证码并原子确认。", "成功进入登录态；失败返回不泄露账号状态的提示。"),
        ("GitHub OAuth 登录", "用户发起 GitHub 授权并携带回调 code。", "校验 state 与回调白名单，换取用户身份，创建或绑定本地账户并签发 HttpOnly JWT。", "跳转工作台；失败进入安全错误页。"),
        ("会话与防滥用", "浏览器携带 Cookie 和 CSRF token 发起登录后请求。", "校验 JWT、账户状态、入口限流和注销状态；验证码超时或超限则拒绝。", "返回当前用户信息、登出结果或限流提示。"),
    ]),
    ("F002", "行情中心（实时行情与 K 线）", [
        ("标的报价", "选择 A 股、美股或加密货币标的。", "按白名单校验市场和 symbol，读取 Java 行情服务的快照并处理数据源异常。", "展示最新价、涨跌幅、时间戳和数据来源状态。"),
        ("K 线与技术指标", "选择周期和数据数量。", "查询 K 线并计算 SMA、EMA、RSI、MACD、支撑阻力等指标。", "返回可视化 K 线和指标摘要，标明研究性结论。"),
    ]),
    ("F003", "AI 智能对话（投研助手）", [
        ("对话请求", "用户输入问题、上下文和可选标的。", "Java 校验登录、功能权限和配额后，以内部令牌调用 Python AI 服务。", "返回结构化回答和金融风险提示。"),
        ("流式输出", "用户发起 SSE 对话并可主动停止。", "Java 代理 Python 流式响应，处理断开取消、超时和错误回退。", "逐段展示文本，结束时记录用量。"),
    ]),
    ("F004", "模拟盘交易", [
        ("模拟下单", "提交标的、买卖方向、数量、杠杆和幂等订单号。", "使用新鲜行情、余额和保证金规则校验；事务内加锁并写入成交记录。", "返回订单、成交和资产变化。"),
        ("持仓与强平", "查看持仓或触发风控扫描。", "根据最新行情计算浮动盈亏和保证金率，超阈值时执行强平。", "展示持仓、风险状态和强平记录。"),
    ]),
    ("F005", "策略回测", [("参数回测", "提交快慢均线、周期和本金。", "Java 读取历史行情，按确定性规则计算收益、回撤和交易明细。", "返回指标卡、净值曲线和可复现结果。")]),
    ("F006", "财报智能解析", [("财报解析", "上传 PDF/图片或粘贴文本。", "提取文本、识别财务与非财务字段，调用 AI 进行结构化复核。", "返回结构化 JSON、关键指标和待核验项。")]),
    ("F007", "智能询报价与价格走势预测", [("报价解读", "选择标的和预测区间。", "汇总行情、供需与财务特征，生成趋势区间并由 AI 解释依据。", "展示现价、区间、趋势要点和风险提示。")]),
    ("F008", "研报情感分析与争议点梳理", [("研报分析", "提交研报文本或文档。", "进行情感分类、观点聚类和论据对比，提取相反观点。", "输出看多/看空/中性标签、摘要和争议点。")]),
    ("F009", "产业链深度挖掘", [("产业链分析", "输入产业链节点或行业关键词。", "组织上下游实体、供需关系和事件传导链，调用 AI 生成解释。", "展示产业链关系、关键厂商和传导路径。")]),
    ("F010", "风险预警模型", [("风险监测", "配置风险指标和阈值。", "计算 VaR、ES 等指标，持续比较阈值并生成风险来源与影响范围。", "展示风险仪表盘、预警等级和应对建议。")]),
    ("F011", "个性化策略生成", [("策略建议", "填写风险偏好和情景假设。", "结合用户画像与研究结果生成展示型策略卡片，并注明非投资建议。", "返回策略说明、适用条件和风险边界。")]),
    ("F012", "市场趋势预测", [("趋势预测", "选择资产类别和时间窗口。", "汇总宏观、行业、企业和舆情特征，使用时序/分类基线输出趋势。", "展示预测方向、依据和不确定性说明。")]),
    ("F013", "运维监控", [("服务健康检查", "监控任务读取服务状态和调用指标。", "检查 Java、Python、数据库和 AI Provider 状态，记录延迟和失败。", "返回健康状态、告警事件和恢复信息。")]),
    ("F014", "管理员账户、AI 配额与功能权限", [
        ("账户管理", "管理员按邮箱、注册方式或状态搜索用户。", "后端校验管理员权限，执行启用/禁用、会话重置和详情查看，并写审计。", "展示账户状态和最近登录信息，不暴露密码或密钥。"),
        ("AI 配额管理", "管理员输入请求次数、Token、周期和调整原因。", "事务内原子调整配额，记录发放、消耗、剩余和周期重置。", "返回调整结果、剩余额度和操作记录。"),
        ("功能权限与审计", "管理员为用户或用户组分配功能权限。", "采用 RBAC 和功能白名单，在接口层强制校验并保存变更前后摘要。", "返回授权结果；越权请求统一拒绝并留痕。"),
    ]),
]


EXTRA_USE_CASES = {
    "策略回测": ("普通用户", ["设置回测参数", "运行双均线回测", "查看收益与回撤"]),
    "财报智能解析": ("普通用户", ["上传财报", "提取财务指标", "查看结构化结果"]),
    "智能询报价与价格走势预测": ("普通用户", ["选择预测标的", "生成报价区间", "查看 AI 解读"]),
    "研报情感分析与争议点梳理": ("普通用户", ["提交研报", "情感分类", "对比争议观点"]),
    "产业链深度挖掘": ("普通用户", ["输入产业链节点", "分析上下游关系", "查看传导路径"]),
    "风险预警模型": ("普通用户", ["配置风险阈值", "计算 VaR 与 ES", "查看预警报告"]),
    "个性化策略生成": ("普通用户", ["填写风险偏好", "生成策略卡片", "查看适用边界"]),
    "市场趋势预测": ("普通用户", ["选择资产类别", "生成趋势预测", "查看预测依据"]),
    "运维监控": ("管理员", ["查看服务健康", "查看 AI 调用指标", "处理告警事件"]),
    "管理员账户、AI 配额与功能权限": ("管理员", ["搜索用户账户", "调整 AI 配额", "分配功能权限", "查看操作审计"]),
}


def write_work_files() -> None:
    WORK.mkdir(parents=True, exist_ok=True)
    UML_DIR.mkdir(parents=True, exist_ok=True)
    source = PRD.read_text(encoding="utf-8")
    def normalize_module_heading(match: re.Match[str]) -> str:
        name = re.sub(r"\s*[｜|]\s*P[012]\s*$", "", match.group(2)).strip()
        name = name.replace("GitHub/邮箱", "GitHub与邮箱")
        return f"### F{int(match.group(1)):03d} {name}"

    source = re.sub(r"^####\s+FR-(\d{2})\s+(.+)$", normalize_module_heading, source, flags=re.MULTILINE)
    SOURCE_MD.write_text(source, encoding="utf-8")
    modules = []
    for module_id, name, funcs in FUNCTIONS:
        sub_functions = []
        for fname, desc, process, output in funcs:
            sub_functions.append({
                "name": fname,
                "desc": desc,
                "input": "用户输入、认证上下文或系统定时任务；仅接受结构完整且符合白名单/额度/权限约束的数据。",
                "process": process,
                "output": output,
            })
        modules.append({"id": module_id, "name": name, "sub_functions": sub_functions})
    FUNCTION_DATA.write_text(json.dumps({"modules": modules}, ensure_ascii=False, indent=2), encoding="utf-8")
    for name, content in DIAGRAMS.items():
        (UML_DIR / name).write_text(content, encoding="utf-8")
    for module_name, (actor, usecases) in EXTRA_USE_CASES.items():
        safe_name = f"用例图_{module_name}.puml"
        lines = [
            "@startuml",
            "left to right direction",
            "skinparam shadowing false",
            "skinparam dpi 140",
            f'actor "{actor}" as U',
            f'rectangle "{module_name}" {{',
        ]
        aliases = []
        for idx, usecase in enumerate(usecases, 1):
            alias = f"UC{idx}"
            aliases.append(alias)
            lines.append(f'  usecase "{usecase}" as {alias}')
        lines.append("}")
        for alias in aliases:
            lines.append(f"U --> {alias}")
        for left, right in zip(aliases, aliases[1:]):
            lines.append(f"{left} -[hidden]down- {right}")
        lines.extend(["@enduml", ""])
        (UML_DIR / safe_name).write_text("\n".join(lines), encoding="utf-8")


def render_diagrams() -> None:
    pngs = []
    for puml in sorted(UML_DIR.glob("*.puml")):
        subprocess.run(["java", "-jar", str(PLANTUML), "-tpng", "-charset", "UTF-8", puml.name], cwd=UML_DIR, check=True)
        pngs.append(puml.with_suffix(".png"))
    expected = [UML_DIR / name.replace(".puml", ".png") for name in DIAGRAMS]
    missing = [str(p) for p in expected if not p.exists()]
    if missing:
        raise RuntimeError("UML 图生成失败: " + ", ".join(missing))


def add_page_break(doc: Document) -> None:
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)


def add_cover_and_toc(doc: Document) -> None:
    body = doc._body._element
    cover = []

    p = doc.add_paragraph()
    # The built-in Word Title style may carry a blue rule; use direct formatting
    # so the cover stays consistent with the formal SRS template.
    p.style = doc.styles["Normal"]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("JARVIS金融投研平台")
    r.bold = True
    r.font.size = Pt(28)
    cover.append(p._p)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Software Requirement Specification")
    r.font.size = Pt(18)
    cover.append(p._p)

    for _ in range(4):
        cover.append(doc.add_paragraph("")._p)

    meta = [
        ("文档版本", "V1.0"),
        ("项目名称", "JARVIS金融投研平台（DeepSeek大模型金融投研系统）"),
        ("小组序号", "01"),
        ("文档日期", "2026年9月7日"),
        ("文档状态", "小组内部评审通过版本"),
    ]
    table = doc.add_table(rows=len(meta), cols=2)
    table.style = "Table Grid"
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for i, (k, v) in enumerate(meta):
        table.cell(i, 0).text = k
        table.cell(i, 1).text = v
        for cell in table.rows[i].cells:
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    run.font.size = Pt(11)
    cover.append(table._tbl)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("本文件仅包含需求、架构和技术选型；任何 API Key、OAuth Secret 或内部令牌均不写入本文档。")
    cover.append(p._p)
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)
    cover.append(p._p)

    toc_title = doc.add_paragraph()
    toc_title.style = doc.styles["Heading 1"]
    toc_title.add_run("目录")
    cover.append(toc_title._p)
    entries = [
        "1 简介",
        "2 总体概述",
        "3 具体需求",
        "4 性能需求",
        "5 接口需求",
        "6 总体设计约束",
        "7 软件质量特性",
        "8 需求分级",
        "9 项目总体架构与技术选型",
    ]
    for entry in entries:
        p = doc.add_paragraph(entry)
        p.paragraph_format.left_indent = Cm(0.8)
        p.paragraph_format.space_after = Pt(4)
        cover.append(p._p)
    p = doc.add_paragraph("注：目录条目与 Heading 样式保持一致；在 Word 中可通过“更新域”刷新页码。")
    p.paragraph_format.space_before = Pt(10)
    cover.append(p._p)
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)
    cover.append(p._p)

    for element in reversed(cover):
        body.insert(0, element)


def shrink_generated_diagrams(doc: Document) -> None:
    """Keep tall generated figures inside one letter-size content area."""
    target_heights_cm = {
        "流程图_主业务流程.png": 18.6,
    }
    for index, shape in enumerate(doc.inline_shapes):
        rid = shape._inline.graphic.graphicData.pic.blipFill.blip.embed
        related = doc.part.related_parts.get(rid)
        if related is None:
            continue
        # The generator adds the main use-case image first and the main flow
        # image second; relationship part names are generic image1/image2.
        if index == 1:
            ratio = shape.height / shape.width
            shape.height = Cm(target_heights_cm["流程图_主业务流程.png"])
            shape.width = int(shape.height / ratio)


def set_cell_text(cell, text: str, bold: bool = False) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(9.5)


def append_architecture_and_stack(doc: Document) -> None:
    doc.add_page_break()
    doc.add_heading("9 项目总体架构与技术选型", level=1)
    doc.add_heading("9.1 系统总体架构", level=2)
    doc.add_paragraph("系统采用前后端分离和内部 AI 服务隔离架构。浏览器和移动端只访问 Java 业务后端；Java 负责认证、账户、CRUD、行情、交易、回测、风控、审计和统一 API 边界；Python 仅作为内部 AI 与数值计算服务。服务端密钥通过环境变量或部署平台密钥注入，禁止进入前端、文档和版本库。")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(UML_DIR / "架构图_系统总体架构.png"), width=Cm(15.2))
    p = doc.add_paragraph("图 9-1 JARVIS 金融投研平台系统总体架构")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_heading("9.2 产品技术选型", level=2)
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER
    headers = ["层次", "技术选型", "主要职责", "选型理由与约束"]
    for i, h in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], h, True)
    rows = [
        ["前端", "Vue 3 + Vite + ECharts；原生 CSS 设计系统", "登录注册、行情/K线、AI 交互、回测、模拟盘和管理员界面", "与现有 frontend 工程一致；组件轻量、构建快；前端只调用 Java /api，不直连 Python 或第三方服务。"],
        ["后端（业务/CRUD）", "Java 17 + Spring Boot 3.3 + Spring Security + JWT HttpOnly Cookie + Spring Data JPA + PostgreSQL/Flyway（本地 H2）", "认证、RBAC、用户/配额/审计、行情、K线、回测、模拟交易、风控和统一 API 网关", "事务、并发控制和数据一致性边界清晰；生产 schema 由 Flyway 管理；金额、数量、价格使用 BigDecimal。"],
        ["AI 端", "Python 3.10+ + FastAPI + Uvicorn + Pydantic + requests；DeepSeek 兼容 Chat Completions Provider", "AI 推理、流式输出适配、财报/研报/产业链分析、数值计算和数据预处理", "复用现有 backend 内部服务；仅监听内网地址并校验内部服务令牌；模型不负责未经 Python 计算的金融数值。"],
        ["数据与外部服务", "PostgreSQL、H2、GitHub OAuth、Resend、行情数据源、Nginx/Cloudflare", "持久化、第三方身份、邮箱验证、行情采集和 HTTPS 入口", "第三方密钥只放服务端环境变量/密钥管理；回调地址采用白名单；外部行情失败时使用缓存和备用源。"],
    ]
    for row in rows:
        cells = table.add_row().cells
        for i, text in enumerate(row):
            set_cell_text(cells[i], text)

    doc.add_heading("9.3 核心数据流与安全边界", level=2)
    for text in [
        "用户请求：前端通过 HTTPS 访问 Java API，Java 完成 JWT、CSRF、账户状态、RBAC 和功能权限校验。",
        "AI 请求：Java 原子扣减用户配额后，通过内部服务令牌调用 Python；Python 再访问 AI Provider，浏览器永不接触 AI API Key。",
        "行情请求：Java 统一采集和缓存行情，前端仅消费 Java 返回的数据；指标和数值计算可由 Python 或 Java 的确定性服务完成并记录口径。",
        "管理审计：账户启停、配额调整、权限变更和高风险操作记录操作者、目标、时间、来源 IP、变更摘要和原因。",
    ]:
        doc.add_paragraph(text, style="List Bullet")

    doc.add_heading("9.4 版本实施边界", level=2)
    doc.add_paragraph("V1.0 以已评审 PRD 为需求基线。P0 功能优先保证认证、行情、AI 对话、模拟盘和回测闭环；P1 功能覆盖财报解析、询报价、研报分析、产业链、风险预警及管理员账户/配额/权限；P2 功能作为增量迭代。生产部署前必须补齐密钥注入、GitHub OAuth/Resend 配置、PostgreSQL 迁移、HTTPS 反向代理和监控告警。")


def complete_template_placeholders(doc: Document) -> None:
    """Replace legacy template placeholders with project-specific baseline requirements."""
    replacements = {
        "本项目的假设和依赖关系待补充。": (
            "项目假设与依赖：用户通过现代 Web 浏览器访问系统；生产环境提供 HTTPS、PostgreSQL、可用的行情数据源、GitHub OAuth 和 Resend 邮件服务；AI Provider 提供 DeepSeek 兼容接口。第三方服务不可用时，系统必须支持缓存、重试、降级提示和管理员可观测。"
        ),
        "数据字典待补充。": (
            "核心数据实体包括：用户、登录身份、角色/权限、AI 配额、行情标的与 K 线、订单、成交、持仓、回测任务与结果、研报/财报文档、风险预警和审计日志。字段类型、空值规则、唯一约束、索引和脱敏规则在数据库设计阶段固化，并通过 Flyway 迁移版本管理。"
        ),
        "E-R关系图待补充。": (
            "核心关系：用户与角色为多对多；用户与登录身份、AI 配额、订单、持仓、回测任务、研报分析和审计日志为一对多；订单与成交记录为一对多；回测任务与回测结果为一对一。管理操作必须关联操作者、目标对象、变更前后摘要和时间。"
        ),
        "时间性能需求待补充。": (
            "性能目标：标准网络环境下首屏加载时间不超过 3 秒；普通 CRUD API P95 不超过 500 毫秒；行情查询与缓存接口 P95 不超过 1 秒；AI 请求首 token 目标不超过 5 秒并支持流式输出。回测、财报解析等长任务采用异步任务和状态查询，不能阻塞主请求。"
        ),
        "系统开放性需求待补充。": (
            "开放性要求：采用前后端分离与 REST/JSON 接口，Java API 作为统一边界；AI Provider、行情数据源和邮件服务通过适配器封装，允许替换实现；配置通过环境变量或密钥管理注入；接口版本化并保留向后兼容策略。"
        ),
        "系统可用性需求待补充。": (
            "可用性目标：核心 API 月度可用性目标 99.5%；异常返回统一错误码、可读提示和请求追踪标识；AI/行情上游失败时启用缓存、有限重试和降级提示；健康检查覆盖 Java、Python、数据库和外部 Provider；支持无状态实例重启及 PostgreSQL 备份恢复。"
        ),
        "软件接口需求待补充。": (
            "软件接口要求：Java API 统一采用 HTTPS + JSON，认证使用 HttpOnly Cookie JWT 并启用 CSRF 防护；Java 调用 Python 使用内网 HTTP 和服务令牌；GitHub OAuth、Resend 和行情 API 仅由后端适配器接入；统一约定超时、重试、限流、幂等键、分页和审计字段。"
        ),
        "系统应具备良好的界面设计，使用者清晰易用，功能高度集中。": (
            "界面友好性要求：采用响应式布局，覆盖桌面端和移动端；关键状态、行情时间、配额余额、风险等级和 AI 输出来源清晰可见；表单提供格式校验、错误定位、加载态和空态；图表提供图例、单位、时间范围和可读的无障碍文本。"
        ),
        "系统应具备良好的可维护性和可管理性。": (
            "可维护性要求：采用分层、模块化设计，统一日志、异常、配置、审计和数据库迁移；关键请求携带追踪标识并记录耗时、错误和上游依赖；提供健康检查、指标和告警；接口契约、部署说明和变更记录随版本维护，支持灰度、回滚和故障复盘。"
        ),
        "本系统遵循Web开发标准和相关行业规范。": (
            "标准符合性要求：遵循 HTTP、HTTPS、JSON、REST、OAuth 2.0/OIDC 相关通用约定，并遵循 OWASP 常见 Web 安全实践；金融金额、价格、数量采用高精度十进制定点计算；日志和导出数据遵循最小化、脱敏和权限隔离原则。"
        ),
        "技术限制待补充。": (
            "技术限制：生产环境固定使用 Java 17、Spring Boot 3.3、Python 3.10+ 和 PostgreSQL；浏览器不得直连 Python 或第三方 API；AI 数值结果必须可追溯，金额/价格/数量分别使用 BigDecimal/Decimal；所有密钥只通过服务端环境变量或密钥管理注入，不得写入前端、文档、日志或版本库。"
        ),
    }
    for paragraph in doc.paragraphs:
        replacement = replacements.get(paragraph.text)
        if replacement is None:
            continue
        for run in paragraph.runs:
            run.text = ""
        paragraph.add_run(replacement)


def compact_requirement_priority(doc: Document) -> None:
    """Keep the short priority legend together after the requirements table."""
    prefixes = ("重要性分类如下：", "A. 必须的", "B. 重要的", "C. 最好有的")
    for paragraph in doc.paragraphs:
        if not paragraph.text.startswith(prefixes):
            continue
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.0
        for run in paragraph.runs:
            run.font.size = Pt(9.5)


def add_accessibility_metadata(doc: Document) -> None:
    """Add descriptive metadata for figures and repeatable table headers."""
    for index, shape in enumerate(doc.inline_shapes, 1):
        doc_pr = shape._inline.docPr
        doc_pr.set("title", f"JARVIS金融投研平台图示{index}")
        doc_pr.set("descr", f"JARVIS金融投研平台需求、流程或架构图示{index}")

    ns = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    for table in doc.tables:
        if not table.rows:
            continue
        tr_pr = table.rows[0]._tr.get_or_add_trPr()
        header = tr_pr.find(f"{{{ns}}}tblHeader")
        if header is None:
            header = OxmlElement("w:tblHeader")
            tr_pr.append(header)


def main() -> None:
    write_work_files()
    render_diagrams()
    subprocess.run([
        sys.executable, str(GENERATOR),
        "--input", str(SOURCE_MD),
        "--uml-dir", str(UML_DIR),
        "--output", str(OUT),
        "--project-name", "JARVIS金融投研平台（DeepSeek大模型金融投研系统）",
        "--group-number", "01",
        "--styles", str(STYLES),
        "--function-data", str(FUNCTION_DATA),
    ], check=True)
    doc = Document(str(OUT))
    shrink_generated_diagrams(doc)
    add_cover_and_toc(doc)
    complete_template_placeholders(doc)
    compact_requirement_priority(doc)
    append_architecture_and_stack(doc)
    add_accessibility_metadata(doc)
    doc.save(str(OUT))
    print(f"SRS generated: {OUT}")
    print(f"UML directory: {UML_DIR}")
    print(f"UML PNG count: {len(list(UML_DIR.glob('*.png')))}")


if __name__ == "__main__":
    main()
